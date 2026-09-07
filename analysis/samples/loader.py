"""
analysis/samples/loader.py
──────────────────────────
Data loading utilities for clinical datasets:
- Discharge summaries: anotated_dataset_v2.xlsx
- Pathology reports: PRAD_reports_eng.xlsx
- Radiology reports: MR-RATE batch00_reports.csv
"""

import os
from pathlib import Path
from typing import Dict, List, Any
try:
    from ..config import DATASETS_DIR
except (ImportError, ValueError):
    from analysis.config import DATASETS_DIR

def load_discharge_summary_rows(indices: List[int]) -> List[Dict[str, Any]]:
    """Load specific rows from annotated discharge summaries Excel."""
    excel_path = DATASETS_DIR / "discharge_summaries" / "anotated_dataset_v2.xlsx"
    if not excel_path.exists():
        raise FileNotFoundError(f"Discharge dataset not found at {excel_path}")

    import openpyxl
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb["annotated_dataset"]

    results = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if idx in indices:
            orig_text = str(row[0]) if row[0] else ""
            diag_icd = str(row[1]) if len(row) > 1 and row[1] else ""
            outcome = str(row[2]) if len(row) > 2 and row[2] else ""
            translated_text = str(row[3]) if len(row) > 3 and row[3] else ""
            trans_diag = str(row[4]) if len(row) > 4 and row[4] else ""
            trans_outcome = str(row[5]) if len(row) > 5 and row[5] else ""
            stay = str(row[6]) if len(row) > 6 and row[6] else ""
            specialty = str(row[7]) if len(row) > 7 and row[7] else ""

            # Use translated English summary if available, otherwise original
            text = translated_text.strip() if translated_text else orig_text.strip()

            results.append({
                "row_index": idx,
                "text": text,
                "icd": trans_diag.strip() if trans_diag else diag_icd.strip(),
                "outcome": trans_outcome.strip() if trans_outcome else outcome.strip(),
                "stay": stay.strip(),
                "specialty": specialty.strip(),
                "datatype": "discharge",
            })
    return results

def load_pathology_report_rows(indices: List[int]) -> List[Dict[str, Any]]:
    """Load specific rows from TCGA PRAD pathology Excel."""
    excel_path = DATASETS_DIR / "pathology_reports" / "TCGA-242" / "PRAD_reports_eng.xlsx"
    if not excel_path.exists():
        raise FileNotFoundError(f"Pathology dataset not found at {excel_path}")

    import openpyxl
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb.active

    results = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if idx in indices:
            specimen_id = str(row[0]) if row[0] else f"PATH_{idx}"
            sub_id = str(row[1]) if len(row) > 1 and row[1] else "1"
            report_text = str(row[2]) if len(row) > 2 and row[2] else ""

            results.append({
                "row_index": idx,
                "specimen_id": specimen_id,
                "sub_id": sub_id,
                "text": report_text.strip(),
                "datatype": "pathology",
            })
    return results

def load_radiology_report_rows(indices: List[int]) -> List[Dict[str, Any]]:
    """Load specific rows from MR-RATE cranial MRI reports."""
    csv_path = DATASETS_DIR / "radiology_reports" / "MR-RATE" / "batch00_reports.csv"
    if not csv_path.exists():
        raise FileNotFoundError(f"Radiology dataset not found at {csv_path}")

    import pandas as pd
    df = pd.read_csv(csv_path)

    results = []
    for idx in indices:
        if idx < len(df):
            row = df.iloc[idx]
            report = str(row.get("report", "")).strip()
            findings = str(row.get("findings", "")).strip()
            impression = str(row.get("impression", "")).strip()
            clinical_info = str(row.get("clinical_information", "")).strip()

            text = report if report else f"Findings:\n{findings}\n\nImpression:\n{impression}"

            results.append({
                "row_index": idx,
                "study_uid": str(row.get("study_uid", f"MR_{idx}")),
                "text": text,
                "clinical_information": clinical_info,
                "findings": findings,
                "impression": impression,
                "datatype": "radiology",
            })
    return results
